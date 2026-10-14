"""
rag.py - Arabic RAG core (Ollama + command-r7b-arabic + bge-m3), hybrid search, ID lookup tool.

Commands:
  python rag.py ingest [source ...]  build the index. A source is a folder, a single file, or a database URL:
                                       folder/file: .xlsx .xlsm .xls .ods .csv .tsv .json .jsonl .md .txt .html .pdf .docx
                                                    .db .sqlite .sqlite3 (SQLite: every table and view)
                                       database:    postgresql://user:pass@host/db, mysql+pymysql://..., mssql+pyodbc://...
                                                    (needs: pip install sqlalchemy + the driver). Or set DB_URL.
  python rag.py ask "سؤالك"          ask one question in the terminal
  python rag.py eval [data_dir]      run the test questions (eval_questions.json or an Excel 'eval...' sheet)
  python rag.py models               list the chat models installed in Ollama

Model: any Ollama chat model works (7B/8B ones like command-r7b-arabic, qwen2.5:7b, llama3.1:8b, aya-expanse:8b).
       LLM_MODEL=qwen2.5:7b   pick the model. If it isn't installed, an installed chat model is used instead.
       NUM_CTX=4096           context window. The prompt is trimmed to fit, so the rules in the system
                              prompt are never cut off. Raise it (8192) if you have the RAM/VRAM.

Excel: every sheet is read as a table (header row found automatically, IDs and dates cleaned).
       Sheets named README / notes / eval... are not indexed. A cell longer than 400 characters
       (or a multi-line cell over 150) is treated as a document and chunked, so a sheet of
       procedure texts works too.
       Put the Excel file alone in its folder, not next to CSV copies of the same data.
UI:   streamlit run ui.py

Offline dev mode (no Ollama needed, for UI work):  set RAG_FAKE=1
"""
import csv, glob, hashlib, json, math, os, pickle, re, sys
from collections import Counter

import numpy as np

LLM_MODEL = os.getenv("LLM_MODEL", "command-r7b-arabic")
EMB_MODEL = os.getenv("EMB_MODEL", "bge-m3")
DATA_DIR = os.getenv("DATA_DIR", "data_excel")
INDEX_PATH = os.getenv("INDEX_PATH", "index.pkl")
TOP_K = int(os.getenv("TOP_K", "5"))
FAKE = os.getenv("RAG_FAKE") == "1"
# "fast" (default) = keyword search only: ingest takes seconds, and Ollama never has to swap
#                    the embedding model and the chat model in and out of memory on each question.
# "hybrid"         = keyword + bge-m3 semantic search: better for loosely worded questions, much slower on CPU.
SEARCH_MODE = os.getenv("SEARCH_MODE", "fast")
# Set NUM_GPU=0 to run the model on the CPU only (fixes "CUDA error" crashes on some laptops; slower)
NUM_GPU = os.getenv("NUM_GPU")
NUM_CTX = int(os.getenv("NUM_CTX", "4096"))
NUM_PREDICT = int(os.getenv("NUM_PREDICT", "400"))
DB_URL = os.getenv("DB_URL")  # optional: also index this database on ingest
MAX_DB_ROWS = int(os.getenv("MAX_DB_ROWS", "200000"))  # per table, so a huge log table can't eat all the RAM

# ------------------------------------------------------------------ 1. Arabic normalization
_DIACRITICS = re.compile(r"[ؐ-ًؚ-ٰٟۖ-ۭـ]")  # tashkeel + tatweel


def normalize(text: str) -> str:
    """Make Arabic spelling variants match: أإآ->ا, ى->ي, ة->ه, ؤ->و, ئ->ي, strip diacritics, Arabic digits->latin."""
    t = _DIACRITICS.sub("", str(text))
    t = re.sub("[إأآٱ]", "ا", t).replace("ى", "ي").replace("ة", "ه").replace("ؤ", "و").replace("ئ", "ي")
    t = t.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    return t.lower()


STOP = set(normalize(w) for w in """في من على الى إلى عن ما ماذا متى كم هل هو هي هذا هذه ان أن او أو و ثم
شو قديش وين ليش كيف بدي اللي الذي التي يا مع كل بعد قبل لا لم لن""".split())


_PREFIXES = ("وال", "بال", "كال", "فال", "لل", "ال")  # one-letter prefixes are handled in tokenize()
_SUFFIXES = ("هما", "تين", "تها", "ها", "هم", "كم", "نا", "تي", "ات", "ون", "ين", "يه", "ه", "ي", "ك")


def stem(w):
    """Light Arabic stemming so هويتي / الهوية / هويه all become the same search word."""
    if w.isdigit() or len(w) < 4:
        return w
    for p in _PREFIXES:
        if w.startswith(p) and len(w) - len(p) >= 3:
            w = w[len(p):]
            break
    for s in _SUFFIXES:
        if w.endswith(s) and len(w) - len(s) >= 3:
            w = w[:-len(s)]
            break
    return w


# Everyday / dialect words people type -> the words used in the official documents
SYNONYMS = {
    "هويه": "بطاقه احوال", "هويتي": "بطاقه احوال", "الهويه": "بطاقه احوال", "بطاقتي": "بطاقه",
    "ضاعت": "فاقد فقدان", "ضايعه": "فاقد", "فقدت": "فاقد فقدان", "انسرقت": "فاقد", "خربت": "تالف", "انكسرت": "تالف",
    "مولود": "ولاده تسجيل", "مواليد": "ولاده", "ابني": "ولاده", "بنتي": "ولاده", "طفل": "ولاده",
    "مات": "وفاه", "توفي": "وفاه", "متوفي": "وفاه", "وفاته": "وفاه",
    "اوراق": "الوثائق المطلوبه", "الاوراق": "الوثائق المطلوبه", "ورق": "الوثائق",
    "قديش": "رسوم", "سعر": "رسوم", "تكلفه": "رسوم", "بكم": "رسوم", "ادفع": "رسوم",
    "دوام": "ساعات الدوام", "بفتحو": "ساعات الدوام", "يفتح": "ساعات الدوام",
    "تزوجت": "زواج", "اتجوز": "زواج", "اطلق": "طلاق", "طلقت": "طلاق",
    "اسمي": "تغيير الاسم", "اغير": "تغيير تصحيح",
    "اجدد": "تجديد", "جدد": "تجديد", "يجدد": "تجديد", "بجدد": "تجديد", "منتهيه": "تجديد انتهاء", "خلصت": "تجديد انتهاء",
    "اطلع": "اصدار", "اسجل": "تسجيل", "بسجل": "تسجيل",
}


def tokenize(text, expand=False):
    words = re.findall(r"[\w]+", normalize(text))
    if expand:  # only for questions: add the official wording for dialect words
        words += [x for w in words for x in SYNONYMS.get(w, "").split()]
    out = []
    for w in words:
        if w in STOP or len(w) < 2:
            continue
        out.append(stem(w))
        # و ب ل ف ك glued to the front (لبطاقة, وبدي): also index the word without it.
        # Done as an extra token, not a replacement, so real words like فاقد / بطاقة stay intact.
        if len(w) >= 5 and w[0] in "وبلفك" and not w.startswith("ال"):
            out.append(stem(w[1:]))
    return out


# ------------------------------------------------------------------ 2. Loading + chunking
def split_text(text, max_chars):
    """Cut long text into pieces of at most max_chars, at a paragraph / line / word break, never mid-word."""
    out = []
    while len(text) > max_chars:
        cut = max(text.rfind("\n\n", 0, max_chars), text.rfind("\n", 0, max_chars), text.rfind(" ", 0, max_chars))
        cut = cut if cut > max_chars // 3 else max_chars
        out.append(text[:cut].strip())
        text = text[cut:].strip()
    if text.strip():
        out.append(text.strip())
    return [p for p in out if p]


def chunk_markdown(text, source, max_chars=1200):
    """Split on headings so each chunk is one topic; keep the doc title on every chunk for context."""
    title = next((l.lstrip("# ").strip() for l in text.splitlines() if l.startswith("# ")), source)
    parts = re.split(r"\n(?=##\s)", text)
    chunks = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        body = p if p.startswith("# ") else f"{title}\n{p}"
        for piece in split_text(body, max_chars):  # hard cap for very long sections
            chunks.append({"text": piece, "source": source, "kind": "doc"})
    return chunks


def row_to_text(row: dict) -> str:
    """Generic: works for ANY table they hand you. 'column: value' lines read well for LLMs."""
    return "\n".join(f"{k}: {v}" for k, v in row.items() if str(v).strip())


def read_text_file(path):
    """Arabic files from government systems are often Windows-1256, not UTF-8. Try both."""
    for enc in ("utf-8-sig", "cp1256"):
        try:
            with open(path, encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def load_table_rows(path):
    """CSV / TSV / JSON / JSONL -> list of row dicts, or None if the file isn't a table."""
    import io
    ext = path.lower().rsplit(".", 1)[-1]
    text = read_text_file(path)
    if ext in ("csv", "tsv"):
        delim = "\t"
        if ext == "csv":  # Excel in Arabic/European locales saves CSV with ; instead of ,
            try:
                delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
            except csv.Error:
                delim = ","
        rows = list(csv.DictReader(io.StringIO(text), delimiter=delim))
        return [{str(k or f"column_{i+1}").strip(): cell_to_str(v) for i, (k, v) in enumerate(r.items())} for r in rows]
    if ext == "jsonl":
        data = [json.loads(l) for l in text.splitlines() if l.strip()]
    else:
        data = json.loads(text)
    if isinstance(data, dict):  # {"records": [...]} / {"data": [...]}: use the first list of objects inside
        data = next((v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)), data)
    if not (isinstance(data, list) and data and isinstance(data[0], dict)):
        return None
    return [{str(k): cell_to_str(v) for k, v in r.items()} for r in data if isinstance(r, dict)]


# ---- Excel ------------------------------------------------------------------------------
SKIP_SHEETS = ("readme", "read me", "notes", "about", "تعليمات", "ملاحظات")
LONG_TEXT = 400  # a cell longer than this is a document (procedure text), not a field


def cell_to_str(v):
    """Excel / databases store IDs as numbers and dates as datetimes. Turn them back into clean text."""
    import datetime
    if v is None or isinstance(v, (bytes, bytearray, memoryview)):  # BLOBs (images, files) can't be searched
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float):
        if v != v:  # NaN from pandas = empty cell
            return ""
        if v.is_integer():
            return str(int(v))  # 9040250165.0 -> "9040250165", never "9.04e9"
    if isinstance(v, (dict, list)):  # nested JSON / JSON database columns
        return json.dumps(v, ensure_ascii=False)
    return str(v).strip()


def read_excel_sheets(path):
    """Yield (sheet_name, rows as dicts) for every sheet that holds a table."""
    if path.lower().endswith((".xls", ".ods")):
        import pandas as pd  # old .xls needs pandas + xlrd, .ods needs pandas + odfpy
        for sheet, df in pd.read_excel(path, sheet_name=None, header=None, dtype=object).items():
            yield sheet, table_from_grid(df.where(df.notna(), None).values.tolist())
        return
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)  # data_only: formula results, not formulas
    for ws in wb.worksheets:
        yield ws.title, table_from_grid(list(ws.iter_rows(values_only=True)))
    wb.close()


def table_from_grid(grid):
    """Find the header row (first row in the top 10 with 2+ filled cells), then build row dicts."""
    grid = [[cell_to_str(c) for c in row] for row in grid]
    grid = [r for r in grid if any(r)]
    if not grid:
        return []
    h_idx = next((i for i, r in enumerate(grid[:10]) if sum(1 for c in r if c) >= 2), 0)
    headers, seen = [], Counter()
    for i, h in enumerate(grid[h_idx]):
        h = re.sub(r"\s+", " ", h) or f"column_{i+1}"
        seen[h] += 1
        headers.append(h if seen[h] == 1 else f"{h}_{seen[h]}")  # duplicate header names
    rows = []
    for r in grid[h_idx + 1:]:
        row = {headers[i]: (r[i] if i < len(r) else "") for i in range(len(headers))}
        if any(row.values()):
            rows.append(row)
    return rows


def rows_to_chunks(rows, source):
    """Any table (Excel sheet, CSV, JSON, database table) -> (chunks, records).
    A row with a long text cell is a document and gets chunked; any other row is one record."""
    chunks, records = [], []
    for i, r in enumerate(rows):
        long_cols = [k for k, v in r.items() if len(v) > LONG_TEXT or (len(v) > 150 and "\n" in v)]
        if long_cols:  # e.g. a 'content' column with a whole procedure: chunk it like a document
            short = [v for k, v in r.items() if k not in long_cols and v]
            label = " | ".join(short)
            doc_source = f"{source} › {short[0]}" if short else f"{source} › {i+1}"  # cite the doc, not just the table
            for k in long_cols:
                for c in chunk_markdown(r[k], doc_source):
                    c["text"] = f"{label}\n{c['text']}" if label else c["text"]
                    chunks.append(c)
        else:
            chunks.append({"text": f"[{source} سجل {i+1}]\n" + row_to_text(r), "source": source, "kind": "row"})
            records.append({"source": source, **r})
    return chunks, records


def load_excel(path, name):
    chunks, records = [], []
    for sheet, rows in read_excel_sheets(path):
        s = sheet.strip().lower()
        if not rows or s in SKIP_SHEETS or s.startswith("eval"):
            continue
        c, r = rows_to_chunks(rows, f"{name} › {sheet}")
        chunks += c
        records += r
        print(f"  {name} › {sheet}: {len(rows)} rows")
    return chunks, records


# ---- Databases ---------------------------------------------------------------------------
def hide_password(url):
    return re.sub(r"(://[^:/@]+):[^@]*@", r"\1:***@", url)


def load_sqlite(path, name):
    """Every table and view of a SQLite file. Standard library only, opened read-only."""
    import pathlib, sqlite3
    con = sqlite3.connect(pathlib.Path(path).resolve().as_uri() + "?mode=ro", uri=True)
    chunks, records = [], []
    try:
        tables = [t for (t,) in con.execute(
            "SELECT name FROM sqlite_master WHERE type IN ('table','view') AND name NOT LIKE 'sqlite_%' ORDER BY name")]
        for t in tables:
            cur = con.execute(f'SELECT * FROM "{t.replace(chr(34), chr(34) * 2)}" LIMIT ?', (MAX_DB_ROWS,))
            cols = [d[0] for d in cur.description]
            rows = [r for r in ({c: cell_to_str(v) for c, v in zip(cols, row)} for row in cur) if any(r.values())]
            c, r = rows_to_chunks(rows, f"{name} › {t}")
            chunks += c
            records += r
            print(f"  {name} › {t}: {len(rows)} rows")
    finally:
        con.close()
    return chunks, records


def load_sql_database(url):
    """Any database SQLAlchemy can reach (PostgreSQL, MySQL/MariaDB, SQL Server, Oracle, SQLite...).
    Reads every table and view, at most MAX_DB_ROWS rows each. Read-only: only SELECTs are sent."""
    try:
        from sqlalchemy import MetaData, Table, create_engine, inspect, select
    except ImportError:
        raise RuntimeError("to read a database URL install SQLAlchemy and the driver, e.g. "
                           "pip install sqlalchemy 'psycopg[binary]'   (PostgreSQL)  /  pymysql (MySQL)  /  pyodbc (SQL Server)")
    engine = create_engine(url)
    insp = inspect(engine)
    db_name = engine.url.database or engine.url.host or "database"
    db_name = os.path.basename(str(db_name))
    chunks, records = [], []
    with engine.connect() as con:
        for t in insp.get_table_names() + insp.get_view_names():
            try:
                table = Table(t, MetaData(), autoload_with=engine)
                result = con.execute(select(table).limit(MAX_DB_ROWS))
                cols = list(result.keys())
                rows = [r for r in ({c: cell_to_str(v) for c, v in zip(cols, row)} for row in result) if any(r.values())]
            except Exception as e:
                print(f"!! skipped table {t}: {e}")
                continue
            c, r = rows_to_chunks(rows, f"{db_name} › {t}")
            chunks += c
            records += r
            print(f"  {db_name} › {t}: {len(rows)} rows")
    engine.dispose()
    return chunks, records


# ---- PDF / Word / HTML -------------------------------------------------------------------
def load_pdf(path, name):
    """Text of every page, chunked page by page so answers can cite the page number."""
    from pypdf import PdfReader
    pages = [(pg.extract_text() or "").strip() for pg in PdfReader(path).pages]
    if sum(len(p) for p in pages) < 20 * max(1, len(pages)):
        print(f"!! {name}: almost no text found - it is probably a scanned PDF (images). "
              "Run it through OCR first (e.g. ocrmypdf in.pdf out.pdf), then ingest the new file.")
    chunks = []
    for n, text in enumerate(pages, 1):
        for c in chunk_markdown(text, name):
            c["text"] = f"(صفحة {n})\n{c['text']}"
            chunks.append(c)
    return chunks


def load_docx(path, name):
    """Paragraphs as a document, and every table in the Word file as table rows."""
    import docx
    d = docx.Document(path)
    # Word headings -> markdown headings, so chunk_markdown splits by topic
    lines = [("## " if p.style is not None and p.style.name.lower().startswith("heading") else "") + p.text
             for p in d.paragraphs if p.text.strip()]
    chunks, records = chunk_markdown("\n".join(lines), name), []
    for i, t in enumerate(d.tables, 1):
        grid = [[cell.text for cell in row.cells] for row in t.rows]
        c, r = rows_to_chunks(table_from_grid(grid), f"{name} › جدول {i}")
        chunks += c
        records += r
    return chunks, records


def html_to_text(html):
    from html.parser import HTMLParser

    class P(HTMLParser):
        def __init__(self):
            super().__init__()
            self.out, self.skip = [], 0

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style"):
                self.skip += 1
            elif tag in ("h1", "h2", "h3"):
                self.out.append("\n## ")
            elif tag in ("p", "br", "div", "li", "tr"):
                self.out.append("\n")

        def handle_endtag(self, tag):
            if tag in ("script", "style"):
                self.skip -= 1

        def handle_data(self, data):
            if not self.skip:
                self.out.append(data)

    p = P()
    p.feed(html)
    return re.sub(r"\n\s*\n+", "\n\n", "".join(p.out)).strip()


# ---- One loader for everything -----------------------------------------------------------
TABLE_EXT = ("csv", "tsv", "json", "jsonl")
EXCEL_EXT = ("xlsx", "xlsm", "xls", "ods")
DB_EXT = ("db", "sqlite", "sqlite3")


def load_file(path, name):
    """One file -> (chunks, records). Unknown file types return nothing."""
    ext = path.lower().rsplit(".", 1)[-1]
    if ext in ("md", "txt"):
        return chunk_markdown(read_text_file(path), name), []
    if ext in ("html", "htm"):
        return chunk_markdown(html_to_text(read_text_file(path)), name), []
    if ext in EXCEL_EXT:
        return load_excel(path, name)
    if ext in TABLE_EXT:
        if ext == "json" and os.path.exists(path[:-5] + ".csv"):
            return [], []  # same table already loaded from CSV
        rows = load_table_rows(path)
        if rows is None:  # JSON that isn't a table: index it as text
            return chunk_markdown(read_text_file(path), name), []
        print(f"  {name}: {len(rows)} rows")
        return rows_to_chunks(rows, name)
    if ext in DB_EXT:
        return load_sqlite(path, name)
    if ext == "pdf":
        return load_pdf(path, name), []
    if ext == "docx":
        return load_docx(path, name)
    return [], []


def load_documents(data_dir=DATA_DIR):
    """data_dir can be a folder (searched recursively), a single file, or a database URL (has '://')."""
    if "://" in str(data_dir):
        try:
            return load_sql_database(data_dir)
        except Exception as e:
            print(f"!! could not read database {hide_password(data_dir)}: {e}")
            return [], []
    if os.path.isfile(data_dir):
        paths, root = [data_dir], os.path.dirname(data_dir)
    else:
        paths, root = sorted(glob.glob(os.path.join(data_dir, "**", "*"), recursive=True)), data_dir
    chunks, records = [], []
    for path in paths:
        name = os.path.relpath(path, root)
        base = os.path.basename(path)
        if os.path.isdir(path) or base.startswith("eval_") or base.startswith("~$"):  # ~$ = Excel lock file
            continue
        try:
            c, r = load_file(path, name)
            chunks += c
            records += r
        except Exception as e:
            print(f"!! skipped {name}: {e}")
    return chunks, records


# ------------------------------------------------------------------ 3. Embeddings
def embed(texts, batch=32):
    if FAKE:  # hashed bag-of-words, only for offline UI development
        out = np.zeros((len(texts), 512), dtype=np.float32)
        for i, t in enumerate(texts):
            for w in tokenize(t):
                out[i, int(hashlib.md5(w.encode()).hexdigest(), 16) % 512] += 1
    else:
        import ollama
        vecs = []
        for i in range(0, len(texts), batch):
            vecs += ollama.embed(model=EMB_MODEL, input=texts[i:i + batch])["embeddings"]
            print(f"\r  embedded {min(i + batch, len(texts))}/{len(texts)}", end="", flush=True)
        print()
        out = np.array(vecs, dtype=np.float32)
    out /= (np.linalg.norm(out, axis=1, keepdims=True) + 1e-9)
    return out


# ------------------------------------------------------------------ 4. BM25 (keyword search: names, IDs, exact terms)
class BM25:
    def __init__(self, docs_tokens, k1=1.5, b=0.75):
        self.k1, self.b = k1, b
        self.docs = [Counter(d) for d in docs_tokens]
        self.lens = np.array([len(d) for d in docs_tokens], dtype=np.float32)
        self.avg = float(self.lens.mean()) if len(self.lens) else 1.0
        df = Counter(w for d in self.docs for w in d)
        n = len(self.docs)
        self.idf = {w: math.log(1 + (n - c + 0.5) / (c + 0.5)) for w, c in df.items()}
        self.inv = {}
        for i, d in enumerate(self.docs):
            for w in d:
                self.inv.setdefault(w, []).append(i)

    def scores(self, q_tokens):
        s = np.zeros(len(self.docs), dtype=np.float32)
        for w in set(q_tokens):
            for i in self.inv.get(w, []):
                tf = self.docs[i][w]
                s[i] += self.idf[w] * tf * (self.k1 + 1) / (tf + self.k1 * (1 - self.b + self.b * self.lens[i] / self.avg))
        return s


# ------------------------------------------------------------------ 5. Index build / load
def build_index(data_dir=DATA_DIR, path=INDEX_PATH):
    """data_dir: one source or a list of them (folders, files, database URLs). DB_URL is always added."""
    sources = [data_dir] if isinstance(data_dir, str) else list(data_dir)
    if DB_URL and DB_URL not in sources:
        sources.append(DB_URL)
    chunks, records = [], []
    for src in sources:
        c, r = load_documents(src)
        chunks += c
        records += r
    shown = ", ".join(hide_password(s) for s in sources)  # never write a DB password into the index or the log
    print(f"loaded {len(chunks)} chunks, {len(records)} table rows from {shown}")
    if not chunks:
        print("!! nothing loaded - check the folder name / database URL and that it holds supported files")
        return
    vecs = embed([c["text"] for c in chunks]) if SEARCH_MODE == "hybrid" or FAKE else None
    with open(path, "wb") as f:
        pickle.dump({"chunks": chunks, "records": records, "vecs": vecs, "data_dir": shown,
                     "emb_model": "fake" if FAKE else (EMB_MODEL if vecs is not None else "none (fast mode)")}, f)
    print(f"index saved -> {path}  (search mode: {SEARCH_MODE})")


class Retriever:
    def __init__(self, path=INDEX_PATH):
        with open(path, "rb") as f:
            d = pickle.load(f)
        self.chunks, self.records, self.vecs = d["chunks"], d["records"], d["vecs"]
        self.bm25 = BM25([tokenize(c["text"]) for c in self.chunks])
        self.by_id = {}       # exact lookup on any column that looks like an ID
        self.name_keys = {}   # name pieces (full / first 3 / first+father+family / first+family) -> people
        self.families = {}    # family name -> people
        self.family_display = {}  # normalized family -> name as written in the data (العواملة)
        for r in self.records:
            for k in id_columns(r):
                self.by_id.setdefault(str(r[k]), []).append(r)
            if not is_person(r):
                continue
            t = tuple(name_tokens(name_of(r)))
            keys = {t, (t[0], t[-1])}
            if len(t) >= 3:
                keys |= {t[:3], (t[0], t[1], t[-1])}
            for key in keys:
                self.name_keys.setdefault(key, []).append(r)
            fam = family_of(r)
            if fam:
                self.families.setdefault(fam, []).append(r)
                self.family_display.setdefault(fam, family_label(r))

    # ---- people tools ("agent tools" the answer step calls before asking the model)
    def person(self, pid):
        """The person row whose OWN id is pid (not rows that merely point at pid)."""
        return next((r for r in self.by_id.get(str(pid), []) if is_person(r) and person_id(r) == str(pid)), None)

    def match_names(self, query):
        """People whose name is written in the question. Longest match wins: a full 4-part name beats
        'first father grandfather', which beats 'first + family'. Returns every match, so duplicates show up."""
        q = name_tokens(query)
        for n in range(6, 1, -1):
            found = []
            for i in range(len(q) - n + 1):
                key = tuple(q[i:i + n])
                if n == 2 and key[-1] not in self.families:  # 2 words only count as "first name + family"
                    continue
                found += [r for r in self.name_keys.get(key, []) if r not in found]
            if found:
                return found
        return []

    def match_family(self, query):
        """'كل أفراد عيلة العواملة' -> the family name, if the question asks for a family."""
        q = name_tokens(query)
        fams = [t for t in q if t in self.families]
        if fams and loose_tokens(query) & FAMILY_WORDS:
            return fams[0]
        return None

    def search(self, query, k=TOP_K):
        """Hybrid = semantic (meaning) + BM25 (exact words), merged with Reciprocal Rank Fusion."""
        sparse = self.bm25.scores(tokenize(query, expand=True))
        if self.vecs is None:  # fast mode: keyword search only
            order = [i for i in np.argsort(-sparse)[:200] if sparse[i] > 0]
            docs, per_doc = [], Counter()
            for i in order:  # at most 2 sections per document, so one long document can't fill every slot
                c = self.chunks[i]
                if c["kind"] == "doc" and per_doc[c["source"]] < 2:
                    per_doc[c["source"]] += 1
                    docs.append(c)
            rows = [self.chunks[i] for i in order if self.chunks[i]["kind"] == "row"]
            if docs and rows:  # procedures first; thousands of one-line records would otherwise crowd them out
                n_rows = min(2, len(rows))
                return docs[:k - n_rows] + rows[:n_rows]
            return (docs or rows)[:k]
        dense = self.vecs @ embed([query])[0]
        rrf = np.zeros(len(self.chunks))
        for scores in (dense, sparse):
            for rank, i in enumerate(np.argsort(-scores)[:50]):
                if scores[i] > 0:
                    rrf[i] += 1 / (60 + rank)
        return [self.chunks[i] for i in np.argsort(-rrf)[:k] if rrf[i] > 0]

    def lookup_ids(self, query):
        """'Agent tool': if the question has a national ID / family book no, fetch the record + linked people."""
        hits = []
        for token in re.findall(r"[A-Z]{0,3}\d{6,}", normalize(query).upper()):
            for r in self.by_id.get(token, []):
                if r not in hits:
                    hits.append(r)
        # follow links one hop (spouse, father, mother...) so "who is X's wife" works.
        # Generic: the first ID column of a row is its own ID; any other ID column points at someone else.
        for r in list(hits):
            cols = id_columns(r)
            for k in cols[1:]:
                for linked in self.by_id.get(str(r[k]), []):
                    own = id_columns(linked)
                    if linked not in hits and own and str(linked[own[0]]) == str(r[k]):
                        hits.append(linked)
        return hits[:8]


ID_HINTS = ("id", "رقم", "هوية", "الوطني", "دفتر", "book", "no")


def id_columns(row):
    """Columns whose name hints at an ID AND whose value looks like one (6+ digits, or letters+digits like FB123456)."""
    return [k for k, v in row.items() if k != "source" and any(h in k.lower() for h in ID_HINTS)
            and re.fullmatch(r"[A-Za-z]{0,3}\d{6,}", str(v))]


def record_label(r):
    cols = id_columns(r)
    return f"سجل {r[cols[0]]}" if cols else f"سجل من {r.get('source', '')}"


# ---- people helpers (generic: work on any table with a name column and an ID column)
NAME_SKIP = ("mother", "father", "grand", "first", "family", "office", "الام", "الاب", "الجد", "الاول", "العايله", "مكتب")
FAMILY_WORDS = {"عيله", "عايله", "عيلت", "عشيره", "دار", "ال", "افراد", "اقارب", "كل", "جميع", "اسماء", "قرايب", "family"}
PRETTY = {"full_name": "الاسم", "national_id": "الرقم الوطني", "birth_date": "تاريخ الميلاد", "governorate": "المحافظة",
          "marital_status": "الحالة الاجتماعية", "is_alive": "على قيد الحياة", "gender": "الجنس"}
TABLE_HINTS = ("birth", "ميلاد", "governorate", "محافظ", "marital", "اجتماعي", "alive", "حي")


def name_tokens(text):
    """Name words, normalized so العواملة = عواملة and عبد الله = عبدالله."""
    words = normalize(text).replace("عبد ", "عبد").split()
    return [w[2:] if w.startswith("ال") and len(w) > 4 else w for w in (re.sub(r"[^\w]", "", x) for x in words) if w]


def loose_tokens(text):
    """name_tokens plus each word without a glued-on و ب ل ف ك / لل (للعيلة -> عيله, لمحمد -> محمد)."""
    q = name_tokens(text)
    out = set(q)
    for w in q:
        if w.startswith("لل") and len(w) > 4:
            out.add(w[2:])
        elif len(w) >= 4 and w[0] in "وبلفك":
            out.add(w[1:])
            if w[1:].startswith("ال") and len(w) > 5:
                out.add(w[3:])
    return out


def name_column(r):
    """The full-name column: a header with name/اسم (not mother/father/office...) holding the most words."""
    best, n = None, 1
    for k, v in r.items():
        kn = normalize(k)
        if ("name" in kn or "اسم" in kn) and not any(x in kn for x in NAME_SKIP) and len(str(v).split()) > n:
            best, n = k, len(str(v).split())
    return best


def name_of(r):
    k = name_column(r)
    return str(r[k]) if k else ""


def is_person(r):
    return bool(name_column(r)) and bool(id_columns(r))


def person_id(r):
    cols = id_columns(r)
    return str(r[cols[0]]) if cols else ""


def family_of(r):
    for k, v in r.items():
        kn = normalize(k)
        if ("family" in kn or "عايل" in kn or "عيل" in kn or "عشير" in kn) and "book" not in kn and "دفتر" not in kn and v:
            t = name_tokens(v)
            return t[-1] if t else None
    t = name_tokens(name_of(r))
    return t[-1] if len(t) >= 2 else None


def family_label(r):
    for k, v in r.items():
        kn = normalize(k)
        if ("family" in kn or "عايل" in kn or "عيل" in kn or "عشير" in kn) and "book" not in kn and "دفتر" not in kn and v:
            return str(v)
    return name_of(r).split()[-1] if name_of(r) else ""


def table_row(r):
    """A short row for the on-screen table: name, ID and a few key columns."""
    out = {PRETTY.get(name_column(r), "الاسم"): name_of(r), "الرقم الوطني": person_id(r)}
    for k, v in r.items():
        kn = normalize(k)
        if any(h in kn for h in TABLE_HINTS) and k != name_column(r) and "place" not in kn and "مكان" not in kn:
            out[PRETTY.get(k, k)] = v
    return out


def short_line(r):
    t = table_row(r)
    return " — ".join(f"{k}: {v}" for k, v in t.items() if v)


# ------------------------------------------------------------------ 6. Generation (Nashmi)
BOT_NAME = "نشمي"

WELCOME = """يا هلا ومرحبا، حيّاك الله ونوّرت! 🇯🇴

أنا **نشمي**، مساعدك الذكي بدائرة الأحوال المدنية. اسمي من النشامى، أهل النخوة والفزعة، وهاد بالزبط شغلي: أفزعلك وأخلّص معاملتك بسرعة وبدون ما تستنى بالدور.

بقدر أساعدك بـ:
- رسوم وأوراق الهوية، وتجديدها، وبدل الفاقد والتالف
- تسجيل المواليد والوفيات، الزواج والطلاق، ودفتر العائلة
- البحث عن أي شخص بالرقم الوطني أو بالاسم الكامل
- عرض كل أفراد أي عيلة مسجلة عنا

احكيلي شو بدك وبالطريقة اللي بتريحك، وأبشر بعزّك 🌿"""

SYSTEM_PROMPT = """إنت «نشمي»، المساعد الذكي لدائرة الأحوال المدنية في الأردن.
بتحكي باللهجة الأردنية البيضاء، بأسلوب مؤدب ودافي زي موظف أردني خدوم. استعمل عبارات زي: «أبشر»، «من عيوني»، «تكرم»، «حيّاك الله»، «يا غالي». ما تستعمل ألفاظ قاسية أبداً.

القواعد (مهمة كتير):
- جاوب بس من «المعلومات المتوفرة» تحت. الأرقام والرسوم والتواريخ والأسماء والأرقام الوطنية انقلها حرفياً زي ما هي.
- إذا المعلومة مش موجودة قول: «والله يا غالي هاي المعلومة مش موجودة عندي.» وما تخمّن، وما تستعمل معلوماتك العامة.
- اذكر رقم المصدر بعد المعلومة هيك [1].
- إذا في «تنبيه» إنه في أكتر من شخص بنفس الاسم: احكي للمستخدم بوضوح إنه في أكتر من شخص، واذكرهم كلهم بأسمائهم وأرقامهم الوطنية، واسأله أي واحد بقصد.
- إذا مكتوب إنه في «جدول معروض للمستخدم»: لا تعيد كتابة كل الأسماء، احكي العدد وإنه القائمة كاملة بالجدول تحت.
- «الأشخاص اللي انحكى عنهم» هم اللي سأل عنهم المستخدم قبل بالمحادثة. إذا قال «هو» أو «إلها» أو «أبوه» أو «الأول» افهم إنه بقصد واحد منهم.
- خليك مختصر: من جملة لتلات جمل، أو نقاط قصيرة للأوراق المطلوبة.

مثال على الأسلوب: «أبشر يا غالي! رسوم تجديد الهوية 5 دنانير [1]، وإذا تأخرت في غرامة دينار عن كل شهر [1].»"""

NOT_FOUND = ("والله يا غالي ما لقيت هالمعلومة عندي 🙏\n\n"
             "بقدر أساعدك بالهوية، المواليد والوفيات، الزواج والطلاق، دفتر العائلة، "
             "أو تدوّر على حدا بالرقم الوطني أو بالاسم الكامل.")

GREET = {"مرحبا", "مرحبتين", "هلا", "اهلا", "اهلين", "السلام", "سلام", "صباح", "مسا", "مساء", "هاي", "hi", "hello", "يعطيك"}
THANKS = {"شكرا", "مشكور", "مشكوره", "تسلم", "يسلمو", "يسلموا", "تسلملي", "thanks", "thank"}
WHO = {"مين انت", "شو اسمك", "مين حضرتك", "عرفني عحالك", "عرف عن حالك", "انت مين", "مين نشمي"}
ORDINALS = {"اول": 0, "ثاني": 1, "تاني": 1, "ثالث": 2, "تالت": 2, "رابع": 3, "خامس": 4, "سادس": 5}
REFERS = {"هو", "هي", "هاد", "هذا", "هذه", "هاي", "هاظ", "الشخص", "نفسه", "نفسها", "اله", "الها", "ابوه", "امه", "ابوها",
          "امها", "اولاده", "اولادها", "زوجته", "زوجها", "مرته", "جوزها", "عمره", "عمرها", "هويته", "هويتها", "ميلاده",
          "ميلادها", "اخوه", "اخته", "اخواته", "ابنه", "بنته", "ابنها", "بنتها", "حالته", "حالتها", "عايش", "عايشه", "ساكن"}


def new_memory():
    """Per-conversation memory: who was talked about, who is 'he/she' right now, open choices between duplicates."""
    return {"people": [], "focus": [], "candidates": [], "family": None, "family_label": None, "table": None, "warning": None}


def _remember(memory, rows):
    """Add people to the conversation memory; someone mentioned again moves to the end (most recent)."""
    for r in rows:
        pid = person_id(r)
        if pid:
            memory["people"] = [p for p in memory["people"] if p["id"] != pid] + [{"id": pid, "name": name_of(r)}]
    memory["people"] = memory["people"][-12:]


def _small_talk(question):
    q = normalize(question)
    words = set(re.findall(r"\w+", q))
    if any(w in q for w in WHO):
        return WELCOME
    if len(words) <= 5 and words & THANKS:
        return "العفو يا غالي، هاد واجبنا! 🌹 إذا بدك أي إشي ثاني أنا حاضر."
    filler = {"عليكم", "ورحمه", "الله", "وبركاته", "نشمي", "الخير", "النور", "يا", "كيفك", "كيف", "الحال", "حالك",
              "شلونك", "العافيه", "عليكو", "صباحو", "مساك"}
    if words & GREET and not (words - GREET - filler):  # only a greeting, no question attached
        return "يا هلا وغلا، حيّاك الله! 🌿 أنا نشمي، شو بقدر أخدمك اليوم؟"
    return None


def _ids_in(text):
    return re.findall(r"[A-Z]{0,3}\d{6,}", normalize(text).upper())


def _refers_back(question):
    q = name_tokens(question)
    return any(w in REFERS or (len(w) >= 4 and w.endswith(("ته", "تها", "وه", "ها", "هم"))) for w in q)


def resolve_people(retriever, question, memory):
    """Decide which people the question is about. Returns (records for the model, notes for the model).
    Fills memory['table'] / memory['warning'] for the UI."""
    notes = []
    # 1. a national ID / family book number in the question
    if _ids_in(question):
        records = retriever.lookup_ids(question)
        main = [r for r in records if is_person(r) and person_id(r) in _ids_in(question)]
        if main:
            memory["focus"], memory["candidates"] = [person_id(r) for r in main], []
            _remember(memory, main)
        return records, notes

    # 2. a name written in the question
    found = [r for r in retriever.match_names(question) if is_person(r)]
    if len(found) == 1:
        pid = person_id(found[0])
        memory["focus"], memory["candidates"] = [pid], []
        _remember(memory, found)
        return retriever.lookup_ids(pid), notes
    if len(found) > 1:
        shown = found[:50]
        memory["candidates"], memory["focus"] = [person_id(r) for r in shown], []
        memory["table"] = [table_row(r) for r in shown]
        memory["warning"] = f"في {len(found)} أشخاص بنفس الاسم. تأكد من الرقم الوطني قبل ما تعتمد أي معلومة."
        _remember(memory, shown[:5])
        notes.append(f"تنبيه: في {len(found)} أشخاص مسجلين بنفس الاسم المذكور بالسؤال. "
                     "احكي للمستخدم إنه في أكتر من شخص، واذكرهم كلهم بأرقامهم الوطنية، واسأله أي واحد بقصد "
                     "(بقدر يكتب الرقم الوطني أو يحكي الأول، التاني...). الجدول معروض للمستخدم كمان.\n"
                     + "\n".join(f"{i+1}. {short_line(r)}" for i, r in enumerate(shown[:10])))
        return [], notes

    # 3. a whole family (named in the question, or "the family we talked about" from memory)
    fam = retriever.match_family(question)
    if not fam and memory["family"] and loose_tokens(question) & (FAMILY_WORDS - {"كل", "جميع", "اسماء", "ال", "دار"}):
        fam = memory["family"]
    if fam:
        members = retriever.families[fam]
        alive = [r for r in members if str(r.get("is_alive", "نعم")) != "لا"]
        memory["family"], memory["candidates"], memory["focus"] = fam, [person_id(r) for r in members], []
        memory["family_label"] = retriever.family_display.get(fam, fam)
        memory["table"] = [table_row(r) for r in members]
        notes.append(f"عيلة {retriever.family_display.get(fam, fam)}: عدد الأفراد المسجلين {len(members)} شخص ({len(alive)} منهم على قيد الحياة). "
                     "القائمة الكاملة معروضة للمستخدم بجدول تحت ردّك، فلا تكتب كل الأسماء. أمثلة: "
                     + "، ".join(name_of(r) for r in members[:5]))
        return [], notes

    # 4. "the first one / التاني" after a duplicates or family list
    q = name_tokens(question)
    if memory["candidates"]:
        for w in q:
            if w in ORDINALS and ORDINALS[w] < len(memory["candidates"]):
                pid = memory["candidates"][ORDINALS[w]]
                memory["focus"], memory["candidates"] = [pid], []
                p = retriever.person(pid)
                _remember(memory, [p] if p else [])
                return retriever.lookup_ids(pid), notes

    # 5. a first name of someone we already talked about ("وشو عن رنا؟", "ارجع لمحمد")
    q_loose = loose_tokens(question)
    for p in reversed(memory["people"]):
        first = (name_tokens(p["name"]) or [""])[0]
        if first and first in q_loose:
            memory["focus"] = [p["id"]]
            row = retriever.person(p["id"])
            _remember(memory, [row] if row else [])
            return retriever.lookup_ids(p["id"]), notes

    # 6. "he / his kids / عمره" -> the person in focus
    if memory["focus"] and _refers_back(question):
        records = []
        for pid in memory["focus"]:
            records += [r for r in retriever.lookup_ids(pid) if r not in records]
        return records, notes
    return [], notes


def build_context(chunks, records, notes=(), memory=None):
    parts = [f"[{i+1}] {n}" for i, n in enumerate(notes)]
    for r in records:
        parts.append(f"[{len(parts)+1}] سجل من السجل المدني:\n" + row_to_text({k: v for k, v in r.items() if k != "source"}))
    for c in chunks:
        parts.append(f"[{len(parts)+1}] من {c['source']}:\n{c['text']}")
    text = "\n\n".join(parts)
    if memory and memory["family"]:
        text += f"\n\nآخر عيلة انحكى عنها: {memory.get('family_label') or memory['family']}"
    if memory and memory["people"]:
        text += "\n\nالأشخاص اللي انحكى عنهم بهالمحادثة (من الأقدم للأحدث):\n" + "\n".join(
            f"- {p['name']} (الرقم الوطني {p['id']})" for p in memory["people"])
    return text


def answer(retriever, question, history=None, stream=False, memory=None):
    """memory: a dict from new_memory(), kept by the UI for the whole chat. After the call,
    memory['table'] (rows to show) and memory['warning'] are set for this answer."""
    history = history or []
    memory = memory if memory is not None else new_memory()
    memory["table"], memory["warning"] = None, None

    reply = _small_talk(question)
    if reply:
        return (iter([reply]) if stream else reply), [], ""

    records, notes = resolve_people(retriever, question, memory)
    # people found -> few document results, so they don't distract the model
    chunks = retriever.search(question, k=1 if notes else (2 if records else TOP_K))
    records = records[:8]
    history = [{"role": m["role"], "content": m["content"]} for m in history[-6:]]  # last 3 turns
    chunks, records, history = fit_to_window(chunks, records, notes, memory, history, question)
    context = build_context(chunks, records, notes, memory)
    sources = [record_label(r) for r in records] + [c["source"] for c in chunks]

    if not (records or notes or chunks):
        return (iter([NOT_FOUND]) if stream else NOT_FOUND), sources, context

    # Context goes in the system message (Command R models ground best that way, and every Ollama chat model accepts it)
    messages = [{"role": "system", "content": f"{SYSTEM_PROMPT}\n\nالمعلومات المتوفرة:\n{context}"}]
    messages += history
    messages.append({"role": "user", "content": question})

    if FAKE:
        text = "[وضع تجريبي بدون نموذج] أقرب نتيجة:\n" + context.split("\n\n")[0]
        return (iter([text]) if stream else text), sources, context

    import ollama
    model = chat_model()
    if stream:
        gen = (p["message"]["content"] for p in ollama.chat(model=model, messages=messages, stream=True, options=llm_options(), keep_alive="60m"))
        return gen, sources, context
    out = ollama.chat(model=model, messages=messages, options=llm_options(), keep_alive="60m")["message"]["content"]
    return out, sources, context


# ---- model settings (any 7B/8B Ollama model) --------------------------------------------
# 7B tokenizers spend more tokens on Arabic than on English (often 1 token per 2-3 letters).
# Count 2 letters per token so the estimate errs on the safe side.
CHARS_PER_TOKEN = float(os.getenv("CHARS_PER_TOKEN", "2"))
EMBEDDING_FAMILIES = ("bert", "nomic-bert", "xlm-roberta")
_MODEL = {"name": None}


def approx_tokens(text):
    return int(len(text) / CHARS_PER_TOKEN) + 4


def fit_to_window(chunks, records, notes, memory, history, question):
    """When a prompt is longer than num_ctx, Ollama silently cuts it from the START, which is the system
    prompt with the rules - a 7B model then answers without them. So trim here instead, least useful
    first: old chat turns, then the lowest-ranked search results, then extra records."""
    chunks, records, history = list(chunks), list(records), list(history)
    budget = NUM_CTX - NUM_PREDICT - 64

    def size():
        ctx = build_context(chunks, records, notes, memory)
        return approx_tokens(SYSTEM_PROMPT) + approx_tokens(ctx) + approx_tokens(question) \
            + sum(approx_tokens(m["content"]) for m in history)

    while size() > budget:
        if history:
            history = history[2:]  # drop the oldest question + answer
        elif len(chunks) > 1:
            chunks.pop()
        elif len(records) > 1:
            records.pop()
        elif chunks and records:
            chunks.pop()
        else:  # one huge chunk or record: shorten its text
            items = chunks or records
            if not items:
                break
            over = int((size() - budget) * CHARS_PER_TOKEN) + 20
            if chunks:
                c = chunks[0]
                if len(c["text"]) <= over:
                    break
                chunks[0] = {**c, "text": c["text"][:len(c["text"]) - over]}
            else:
                r = records[0]
                text = row_to_text({k: v for k, v in r.items() if k != "source"})
                if len(text) <= over:
                    break
                records[0] = {"source": r.get("source", ""), "سجل": text[:len(text) - over]}
    return chunks, records, history


def llm_options():
    # smaller context window + capped answer length = much faster on a laptop
    opts = {"temperature": 0.2, "num_ctx": NUM_CTX, "num_predict": NUM_PREDICT}
    if NUM_GPU is not None:
        opts["num_gpu"] = int(NUM_GPU)
    return opts


def installed_chat_models():
    """Chat models pulled in Ollama (embedding models like bge-m3 are left out)."""
    import ollama
    out = []
    for m in ollama.list().models:
        family = (m.details.family if m.details else "") or ""
        if family in EMBEDDING_FAMILIES or "embed" in m.model:
            continue
        out.append((m.model, m.details.parameter_size if m.details else ""))
    return out


def chat_model():
    """LLM_MODEL if it is installed, otherwise the installed chat model closest to 7B, so the app
    still starts after switching machines or models. Raises a clear error if there is none."""
    if _MODEL["name"]:
        return _MODEL["name"]
    models = installed_chat_models()
    names = [n for n, _ in models]
    want = LLM_MODEL if ":" in LLM_MODEL else LLM_MODEL + ":latest"
    if want in names or LLM_MODEL in names:
        _MODEL["name"] = LLM_MODEL
    elif models:
        def dist_to_7b(m):
            size = re.match(r"([\d.]+)\s*([BM])", m[1] or "", re.I)
            billions = float(size.group(1)) / (1000 if size.group(2).upper() == "M" else 1) if size else 99
            return abs(billions - 7)
        _MODEL["name"] = min(models, key=dist_to_7b)[0]
        print(f"!! model '{LLM_MODEL}' is not installed - using '{_MODEL['name']}' instead "
              f"(install it with: ollama pull {LLM_MODEL})")
    else:
        raise RuntimeError(f"no chat model installed in Ollama - run: ollama pull {LLM_MODEL}")
    return _MODEL["name"]


def warmup():
    """Load the chat model into memory once, so the first real question isn't the slow one."""
    if FAKE:
        return
    import ollama
    ollama.chat(model=chat_model(), messages=[{"role": "user", "content": "مرحبا"}],
                options={**llm_options(), "num_predict": 1}, keep_alive="60m")


# ------------------------------------------------------------------ 7. CLI + eval
def load_eval(data_dir=DATA_DIR):
    """Test questions from eval_questions.json, or from any Excel sheet whose name starts with 'eval'.
    Excel columns: question, expected_answer, source (optional) - or just the first two columns."""
    path = os.path.join(data_dir, "eval_questions.json")
    if os.path.exists(path):
        return json.load(open(path, encoding="utf-8"))
    qs = []
    for x in glob.glob(os.path.join(data_dir, "**", "*.xls*"), recursive=True):
        for sheet, rows in read_excel_sheets(x):
            if not sheet.strip().lower().startswith("eval"):
                continue
            for r in rows:
                vals = list(r.values())
                q = r.get("question", vals[0] if vals else "")
                a = r.get("expected_answer", vals[1] if len(vals) > 1 else "")
                if q:
                    qs.append({"q": q, "a": a, "source": r.get("source", "")})
    return qs


def run_eval(retriever, data_dir=DATA_DIR):
    qs = load_eval(data_dir)
    if not qs:
        print("no test questions found: add data/eval_questions.json or an Excel sheet named eval_questions")
        return
    hit = 0
    for q in qs:
        out, sources, ctx = answer(retriever, q["q"])
        src, ans = str(q.get("source", "")), str(q.get("a", ""))
        expected_in_ctx = src == "none" or (src and any(s.startswith(src) or src in s for s in sources)) \
            or (ans and normalize(ans)[:12] in normalize(ctx))
        hit += expected_in_ctx
        print(f"\nQ: {q['q']}\n  expected: {q['a']}\n  got: {str(out)[:300]}\n  retrieval {'OK' if expected_in_ctx else 'MISS'}")
    print(f"\nretrieval hit-rate: {hit}/{len(qs)}  (read the answers yourself to judge correctness)")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "help"
    if cmd == "ingest":
        build_index(sys.argv[2:] or DATA_DIR)
    elif cmd == "ask":
        out, src, _ = answer(Retriever(), " ".join(sys.argv[2:]))
        print(out, "\n\nsources:", src)
    elif cmd == "eval":
        run_eval(Retriever(), sys.argv[2] if len(sys.argv) > 2 else DATA_DIR)
    elif cmd == "models":
        for n, size in installed_chat_models():
            print(f"  {n:40} {size}{'   <- LLM_MODEL' if n.split(':')[0] == LLM_MODEL.split(':')[0] else ''}")
    else:
        print(__doc__)
